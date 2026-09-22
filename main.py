import os
import re
import json
import requests
import sys
from datetime import datetime, timedelta
from copy import deepcopy
from bs4 import BeautifulSoup
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from typing import List, Optional
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

class TerminalColors:
    """Provides colors only when the terminal supports ANSI escape codes."""
    RESET = "\033[0m"
    CYAN = "\033[96m"
    GREEN = "\033[92m"
    YELLOW = "\033[93m"
    RED = "\033[91m"
    DIM = "\033[2m"

    @classmethod
    def enabled(cls) -> bool:
        return sys.stdout.isatty() and "NO_COLOR" not in os.environ

    @classmethod
    def paint(cls, text: str, color: str) -> str:
        if not cls.enabled():
            return text
        return f"{color}{text}{cls.RESET}"

def color_print(text: str, color: str = ""):
    """Prints terminal text with optional ANSI color."""
    print(TerminalColors.paint(text, color) if color else text)

def require_api_key() -> None:
    """Ensures Gemini credentials are configured before an API request is made."""
    if not (os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")):
        raise ValueError("GEMINI_API_KEY or GOOGLE_API_KEY environment variable not found!")

def create_gemini_client() -> genai.Client:
    """Creates a Gemini client with explicit timeout and retry behavior."""
    require_api_key()
    return genai.Client(
        http_options=types.HttpOptions(
            timeout=30000,
            retry_options=types.HttpRetryOptions(attempts=1),
        ),
    )

class GeminiRequestError(RuntimeError):
    """Indicates that a required Gemini request did not complete successfully."""

# --- Structured Pydantic Schemas ---

class BulletChange(BaseModel):
    section_name: str = Field(description="Section where the change occurred (e.g., Professional Summary, Experience, Skills).")
    job_requirement: str = Field(description="The requirement/keyword identified in the job description.")
    matched_candidate_experience: str = Field(description="Candidate's real background/skill matched to this requirement.")
    original_text: str = Field(description="Exact snippet or full sentence from the Master Resume to replace.")
    replacement_text: str = Field(description="The tailored replacement text (plain text, NO markdown asterisks).")

class EmploymentHistoryChange(BaseModel):
    employer: str = Field(description="The employer whose employment-history section is being changed.")
    role_context: str = Field(description="The exact role heading from the Master Resume, used only to target the correct role subsection.")
    job_requirement: str = Field(description="The relevant requirement or keyword from the target job description.")
    matched_candidate_experience: str = Field(description="The exact verified experience that supports this change.")
    action_type: str = Field(description="One of: Reframed, Reordered, or Newly Added.")
    original_text: str = Field(default="", description="Exact existing text to replace, required for Reframed changes.")
    replacement_text: str = Field(default="", description="Replacement text, required for Reframed changes.")
    anchor_text: str = Field(default="", description="Exact existing paragraph after which a new bullet should be inserted, required for Newly Added changes.")
    new_bullet: str = Field(default="", description="New employment-history bullet, required for Newly Added changes.")
    remove_text: str = Field(default="", description="Exact existing bullet to remove when adding a bullet to a full employer section.")

class ResumeTailorResponse(BaseModel):
    changes: List[BulletChange]
    employment_history_changes: List[EmploymentHistoryChange] = Field(default_factory=list)

class PriorityItem(BaseModel):
    priority_rank: int = Field(description="Priority rank based on job description hierarchy (1 = highest importance).")
    job_requirement_matched: str = Field(description="The top requirement or skill from the job description matched here.")
    content_text: str = Field(description="The text for this bullet point or skill entry (plain text, NO markdown).")
    action_type: str = Field(description="'Reordered', 'Reframed', or 'Newly Added Experience'.")
    original_text: str = Field(default="", description="Exact existing paragraph to replace for a Reframed item; empty for other items.")
    anchor_text: str = Field(default="", description="Exact existing paragraph after which to insert content_text for a Newly Added item; empty for other items.")
    employer: str = Field(default="", description="Employer for an employment-history item; empty for Skills or other sections.")
    role_context: str = Field(default="", description="The exact role heading from the Master Resume for an employment-history item.")
    remove_text: str = Field(default="", description="Exact existing bullet to remove when adding an employment bullet to a full employer section.")

class PrioritySection(BaseModel):
    section_title: str = Field(description="Section name (e.g., Core Competencies, Key Skills, Professional Experience).")
    items: List[PriorityItem] = Field(description="List of bullet points or entries, strictly ordered by priority rank.")

class ReorderAndEnhanceResponse(BaseModel):
    priority_sections: List[PrioritySection]

class TailoringResponse(BaseModel):
    company: str = Field(description="Company name extracted from the target job description.")
    title: str = Field(description="Job title extracted from the target job description.")
    changes: List[BulletChange] = Field(default_factory=list)
    employment_history_changes: List[EmploymentHistoryChange] = Field(default_factory=list)
    priority_sections: List[PrioritySection] = Field(default_factory=list)

SYSTEM_PROMPT_PATH = "gemini_system_prompt.txt"

MODEL_FALLBACKS = [
    "gemini-3.7-flash",
    "gemini-3.6-flash",
    "gemini-3.5-flash",
    "gemini-3.5-flash-lite"
]

def get_gemini_error_code(error: Exception) -> int | str:
    """Extracts the API error code even when the SDK does not expose status_code."""
    status_code = getattr(error, "status_code", None)
    if status_code is not None:
        return status_code

    error_text = str(error)
    code_match = re.search(r"['\"]code['\"]\s*:\s*(\d+)|\b(?:HTTP\s+)?(\d{3})\b", error_text, re.IGNORECASE)
    if code_match:
        return int(code_match.group(1) or code_match.group(2))

    status_match = re.search(r"status['\"]?\s*[:=]\s*['\"]([A-Z_]+)", error_text, re.IGNORECASE)
    return status_match.group(1).upper() if status_match else "UNKNOWN"

def report_gemini_error(error: Exception, operation: str, model: str) -> bool:
    """Prints diagnostic details for a failed Gemini request."""
    status_code = get_gemini_error_code(error)
    error_text = str(error)
    is_quota_error = status_code == 429 or "RESOURCE_EXHAUSTED" in error_text or "429" in error_text

    color_print("\n--- Gemini request debug ---", TerminalColors.RED)
    color_print(f"   Operation       : {operation}", TerminalColors.YELLOW)
    color_print(f"   Model           : {model}", TerminalColors.YELLOW)
    color_print(f"   Exception type  : {type(error).__name__}", TerminalColors.RED)
    color_print(f"   Error code      : {status_code}", TerminalColors.RED)
    color_print(f"   Error details   : {error_text}", TerminalColors.DIM)

    if not is_quota_error:
        if status_code == 404 and is_unavailable_model_error(error):
            color_print("   Diagnosis       : The selected model is unavailable or does not support generateContent for this API version.", TerminalColors.YELLOW)
        else:
            color_print("   Diagnosis       : This is not identified as a 429 quota response.", TerminalColors.YELLOW)
        color_print("--- End Gemini request debug ---", TerminalColors.RED)
        return False

    retry_match = re.search(r"retry in\s+([\d.]+)s", error_text, re.IGNORECASE)
    if retry_match is None:
        retry_match = re.search(r"retryDelay.*?(\d+)s", error_text, re.IGNORECASE)

    color_print(f"\n⚠️ Gemini quota limit reached during {operation}.", TerminalColors.RED)
    if "PerDay" in error_text or "per day" in error_text.lower():
        color_print("   This response indicates the free-tier daily request quota was exceeded, not necessarily token depletion.", TerminalColors.YELLOW)
        color_print("   The daily quota must reset or the Gemini project needs a plan/billing change.", TerminalColors.YELLOW)
    else:
        color_print("   This is a temporary rate/quota limit response from Gemini.", TerminalColors.YELLOW)

    if retry_match:
        retry_seconds = float(retry_match.group(1))
        retry_at = datetime.now().astimezone() + timedelta(seconds=retry_seconds)
        color_print(f"   API retry suggestion: wait about {retry_seconds:g} seconds.", TerminalColors.YELLOW)
        color_print(f"   Earliest suggested retry: {retry_at.strftime('%Y-%m-%d %H:%M:%S %Z')}.", TerminalColors.YELLOW)
    else:
        color_print("   Gemini did not provide a retry time in the response.", TerminalColors.YELLOW)
    color_print("--- End Gemini request debug ---", TerminalColors.RED)
    return True

def is_retryable_gemini_error(error: Exception) -> bool:
    """Returns whether another model may reasonably succeed."""
    status_code = getattr(error, "status_code", None)
    error_text = str(error).upper()
    return status_code in (429, 500, 502, 503, 504) or any(
        marker in error_text for marker in (
            "RESOURCE_EXHAUSTED",
            "UNAVAILABLE",
            "TIMEOUT",
            "DEADLINE_EXCEEDED",
            "GATEWAY TIMEOUT",
            "RATE LIMIT",
            "HTTP 500",
            "HTTP 502",
            "HTTP 503",
            "HTTP 504",
        )
    )

def is_unavailable_model_error(error: Exception) -> bool:
    """Returns whether Gemini rejected the selected model itself."""
    error_text = str(error).lower()
    return (
        "model" in error_text
        and ("not found" in error_text or "not supported for generatecontent" in error_text)
    )

def load_employment_history(path: str = "employment_history.json") -> dict:
    """Loads the user-maintained source of truth for employment-history claims."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"Verified employment history missing: {path}")

    with open(path, "r", encoding="utf-8") as file:
        employment_history = json.load(file)

    if not isinstance(employment_history, dict) or not employment_history:
        raise ValueError(f"'{path}' must contain at least one employer entry.")

    return employment_history

def load_system_prompt(path: str = SYSTEM_PROMPT_PATH) -> str:
    """Loads the editable Gemini system prompt from disk."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"Gemini system prompt missing: {path}")

    with open(path, "r", encoding="utf-8") as file:
        prompt = file.read().strip()

    if not prompt:
        raise ValueError(f"'{path}' must contain the Gemini system prompt.")

    return prompt

def get_unique_filepath(base_path: str) -> str:
    """
    Checks if a file exists. If it does, appends _1, _2, etc., before the file extension.
    Example: 'tailored_resume_Company_Role.docx' -> 'tailored_resume_Company_Role_1.docx'
    """
    if not os.path.exists(base_path):
        return base_path

    filename, ext = os.path.splitext(base_path)
    counter = 1
    new_path = f"{filename}_{counter}{ext}"

    while os.path.exists(new_path):
        counter += 1
        new_path = f"{filename}_{counter}{ext}"

    return new_path

def get_job_description(url: str, fallback_file: str = "job_input.txt") -> tuple[str, bool]:
    """Attempts to scrape the job description from a URL, falling back to job_input.txt."""
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept-Language": "en-US,en;q=0.9",
    }
    print(f"\n[1/5] Attempting web scraping for: {url}...")
    
    try:
        response = requests.get(url, headers=headers, timeout=12)
        response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")
        for element in soup(["script", "style", "nav", "footer", "header", "form", "svg"]):
            element.decompose()

        text = soup.get_text(separator="\n")
        lines = (line.strip() for line in text.splitlines())
        clean_text = "\n".join(chunk for line in lines for chunk in line.split("  ") if chunk)

        if len(clean_text) > 300:
            print("✅ SCRAPING SUCCESSFUL!")
            print(f"   Fetched {len(clean_text)} characters from the URL.")
            return clean_text[:8000], True
        else:
            print("❌ SCRAPING FAILED: Page returned minimal content (likely anti-bot blocked).")

    except Exception as e:
        print(f"❌ SCRAPING FAILED: Network/HTTP error occurred ({e}).")

    print("\n" + "="*60)
    print("📋 MANUAL FALLBACK INITIATED")
    if not os.path.exists(fallback_file):
        with open(fallback_file, "w", encoding="utf-8") as f:
            f.write("")
        print(f"-> Created blank '{fallback_file}' in your workspace folder.")

    print(f"1. Open '{fallback_file}' in VS Code and paste the Job Description text.")
    print("2. Save the file (Ctrl+S).")
    input("3. Press ENTER here in the terminal once saved...")
    print("="*60)

    with open(fallback_file, "r", encoding="utf-8") as f:
        manual_text = f.read().strip()

    if not manual_text:
        raise ValueError(f"'{fallback_file}' is empty!")

    print(f"✅ Loaded {len(manual_text)} characters from '{fallback_file}'.")
    return manual_text, False

def get_consolidated_tailoring(
    master_text: str,
    job_desc: str,
    employment_history: dict,
) -> TailoringResponse:
    """Uses one Gemini request to produce all tailoring decisions."""
    print("\n[3/5] Analyzing resume, job description, and employment history with Gemini...")
    client = create_gemini_client()
    system_prompt = load_system_prompt()
    prompt = f"""
### MASTER RESUME TEXT
{master_text}

### TARGET JOB DESCRIPTION
{job_desc}

### VERIFIED EMPLOYMENT HISTORY (SOURCE OF TRUTH)
{json.dumps(employment_history, indent=2)}

Analyze the Master Resume against the Target Job Description and return one structured JSON response containing:
- company: the target company name.
- title: the target job title.
- changes: actionable Skills edits only.
- employment_history_changes: actionable Employment History edits only.
- priority_sections: the ordered application plan for those same changes.

For company and title, return only the extracted value. If either cannot be identified, use 'TargetCompany' or 'Role'.

For changes (Skills only):
- Set section_name to 'Skills'.
- Set original_text to the exact existing Skills entry or text to replace.
- Set replacement_text to the concise improved Skills text.
- Use one primary skill, tool family, platform, or competency per entry. Related tools may be listed with that skill.
- Keep the replacement short. Do not join unrelated skills or add explanatory parentheticals.
- Use only a skill explicitly present in the Master Resume or Verified Employment History, unless the wording is a direct, unambiguous shortening of that same skill.
- Do not invent categories such as 'Container Security' from Docker or Kubernetes.
- Keep separate concepts separate: SIEM is not Incident Investigation; Terraform is not Container Security; Cloud Security is not AWS Security Configurations.
- If an existing entry contains multiple distinct skills, preserve the useful separate skills rather than merging them.
- Do not include employment responsibilities or employer-specific bullets.

Employment History wording:
- Match the wording to the Target Job Description. If it names AWS WAF or Cloudflare, mention the supported named vendor. If it only says WAF or WAF rules, use WAF or WAF rules without listing vendors.
- Keep detailed vendor breadth in Skills unless a vendor is explicitly requested or materially important to the role.
- Include complementary skills such as Terraform only when they form one clear activity with the main responsibility. Do not join unrelated skills merely to include more keywords.
- For a generic WAF requirement, prefer wording such as 'Deploy, manage, and improve application-layer WAF rules using Terraform for Infrastructure as Code' only when that combined activity is supported and reads naturally. Do not force this wording when Terraform is not relevant to the target requirement.

For employment_history_changes (Employment History only):
- Set employer and role_context to identify the original employer subsection and role.
- Set role_context to the exact role heading from the Master Resume. Do not combine separate roles at the same employer.
- For 'Reframed', set original_text to the exact existing bullet and replacement_text to its truthful improved version.
- For 'Newly Added', set anchor_text to an exact existing bullet in the same employer and role subsection, and set new_bullet to the concise supported bullet.
- If the role already has four bullets, set remove_text to the exact least relevant existing bullet to remove before adding the new bullet.
- Never exceed four bullets for any individual role. Keep separate roles at the same employer separate.
- Do not place employment changes in Skills.

For priority_sections (application plan only):
- Include only changes already represented in 'changes' or 'employment_history_changes'.
- Use section_title 'Skills' or 'Employment History - <Employer> - <Role>'.
- For a Skills replacement, include original_text and content_text.
- For an employment reframe, include employer, role_context, original_text, and content_text.
- For an employment addition, include employer, role_context, anchor_text, content_text, and remove_text when required.
- Do not create recommendations that cannot be applied.

Before returning the response:
- Read every Skills entry and every Employment History bullet.
- Compare them with the Target Job Description and Verified Employment History.
- Prefer a truthful reframe or replacement over adding content.
- Add or remove employment bullets only when this improves relevance and preserves the four-bullet maximum.
- Do not duplicate a change across output lists.
- Keep all generated text concise, professional, truthful, and plain text without markdown.
"""

    last_error = None
    for model_index, model in enumerate(MODEL_FALLBACKS):
        color_print(f"   Trying model {model} ({model_index + 1}/{len(MODEL_FALLBACKS)})...", TerminalColors.CYAN)
        try:
            response = client.models.generate_content(
                model=model,
                contents=prompt,
                config=types.GenerateContentConfig(
                    system_instruction=system_prompt,
                    temperature=0.2,
                    response_mime_type="application/json",
                    response_schema=TailoringResponse,
                ),
            )
            result = TailoringResponse.model_validate_json(response.text)
            color_print(f"   Gemini analysis succeeded with {model}.", TerminalColors.GREEN)
            return result
        except Exception as error:
            last_error = error
            report_gemini_error(error, "consolidated tailoring", model)
            retryable = is_retryable_gemini_error(error)
            unavailable_model = is_unavailable_model_error(error)
            has_fallback = model_index < len(MODEL_FALLBACKS) - 1
            should_fallback = (retryable or unavailable_model) and has_fallback
            color_print(
                f"   Fallback decision: retryable={retryable}, "
                f"unavailable_model={unavailable_model}, next_model_available={has_fallback}."
                , TerminalColors.YELLOW
            )
            if not should_fallback:
                break
            color_print(f"   Retrying with fallback model: {MODEL_FALLBACKS[model_index + 1]}.", TerminalColors.CYAN)

    raise GeminiRequestError(
        f"All configured Gemini models failed during consolidated tailoring ({type(last_error).__name__}): "
        f"{last_error}; no document was created."
    ) from last_error

def extract_text_from_docx(docx_path: str) -> str:
    """Extracts plain text from a Word document."""
    print(f"\n[2/5] Reading resume content from '{docx_path}'...")
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Master resume missing: {docx_path}")
    
    doc = Document(docx_path)
    paragraph_tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
    full_text = [
        paragraph.text
        for element in doc.element.body.iter()
        if element.tag == paragraph_tag
        for paragraph in [Paragraph(element, element.getparent())]
        if paragraph.text.strip()
    ]

    print(f"   Loaded {len(full_text)} resume paragraphs, including text-box content.")

    return "\n".join(full_text)

def replace_text_in_docx_inplace(
    original_docx_path: str,
    replacements: List[BulletChange],
    employment_changes: List[EmploymentHistoryChange],
    priority_data: ReorderAndEnhanceResponse,
    output_docx_path: str,
):
    """
    Replaces original text with tailored text directly in paragraphs/tables 
    while preserving original document styles and layout.
    """
    print(f"\n[5/5] Applying text updates to Word document '{output_docx_path}'...")
    doc = Document(original_docx_path)

    applied_count = 0

    def all_paragraphs():
        for element in doc.element.body.iter():
            if element.tag == "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p":
                yield Paragraph(element, element.getparent())

    paragraphs = list(all_paragraphs())

    def paragraph_index(paragraph: Paragraph) -> int:
        return next(index for index, item in enumerate(paragraphs) if item._p is paragraph._p)

    def range_for_heading(start_text: str, end_texts: tuple[str, ...]) -> tuple[int, int]:
        start = next((index for index, paragraph in enumerate(paragraphs)
                      if paragraph.text.strip().lower() == start_text.lower()), None)
        if start is None:
            return 0, len(paragraphs)
        end = next((index for index in range(start + 1, len(paragraphs))
                    if paragraphs[index].text.strip().lower() in end_texts), len(paragraphs))
        return start + 1, end

    employment_start, employment_end = range_for_heading(
        "Employment History",
        ("education and professional training",),
    )

    def employer_range(employer: str) -> tuple[int, int] | None:
        employer_lower = employer.lower()
        current_employment_start, current_employment_end = range_for_heading(
            "Employment History",
            ("education and professional training",),
        )
        start = next((index for index in range(current_employment_start, current_employment_end)
                      if employer_lower in paragraphs[index].text.lower()), None)
        if start is None:
            return None
        next_employer = next((index for index in range(start + 1, current_employment_end)
                              if " at " in paragraphs[index].text.lower() and
                              employer_lower not in paragraphs[index].text.lower()), current_employment_end)
        return start, next_employer

    def role_range(employer: str, role_context: str) -> tuple[int, int] | None:
        employer_bounds = employer_range(employer)
        if employer_bounds is None or not role_context.strip():
            return None

        employer_start, employer_end = employer_bounds
        role_lower = role_context.strip().lower()
        role_start = next((index for index in range(employer_start, employer_end)
                           if role_lower in paragraphs[index].text.lower()), None)
        if role_start is None:
            return None

        role_end = next((index for index in range(role_start + 1, employer_end)
                         if " at " in paragraphs[index].text.lower()), employer_end)
        return role_start, role_end

    def is_bullet(paragraph: Paragraph) -> bool:
        if paragraph._p.pPr is None:
            return False
        return paragraph._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None

    def remove_paragraph(paragraph: Paragraph):
        parent = paragraph._p.getparent()
        parent.remove(paragraph._p)

    def replace_in_range(original_text: str, replacement_text: str, bounds: tuple[int, int]) -> bool:
        start, end = bounds
        for paragraph in paragraphs[start:end]:
            if replace_paragraph_text(paragraph, original_text, replacement_text):
                return True
        return False

    def replace_paragraph_text(paragraph: Paragraph, original_text: str, replacement_text: str) -> bool:
        if original_text not in paragraph.text:
            return False

        runs = paragraph.runs
        if not runs:
            paragraph.add_run(replacement_text)
            return True

        combined_text = "".join(run.text or "" for run in runs)
        if original_text not in combined_text:
            return False

        replacement_start = combined_text.index(original_text)
        replacement_end = replacement_start + len(original_text)
        cursor = 0
        inserted = False
        for run in runs:
            run_text = run.text or ""
            run_start = cursor
            run_end = cursor + len(run_text)
            overlap_start = max(replacement_start, run_start)
            overlap_end = min(replacement_end, run_end)
            if overlap_start < overlap_end:
                before = run_text[:overlap_start - run_start]
                after = run_text[overlap_end - run_start:]
                run.text = before + (replacement_text if not inserted else "") + after
                inserted = True
            cursor = run_end
        return inserted

    def apply_employment_change(change: EmploymentHistoryChange) -> bool:
        bounds = role_range(change.employer, change.role_context)
        if bounds is None:
            return False

        if change.original_text.strip() and change.replacement_text.strip():
            start, end = bounds
            return any(
                replace_paragraph_text(paragraph, change.original_text.strip(), change.replacement_text.strip())
                for paragraph in paragraphs[start:end]
            )

        if not change.anchor_text.strip() or not change.new_bullet.strip():
            return False

        start, end = bounds
        employer_paragraphs = paragraphs[start:end]
        if any(change.new_bullet.strip() == paragraph.text.strip() for paragraph in employer_paragraphs):
            return False

        bullet_count = sum(is_bullet(paragraph) for paragraph in employer_paragraphs if paragraph.text.strip())
        if bullet_count >= 4:
            removable = next((paragraph for paragraph in employer_paragraphs
                              if change.remove_text.strip() in paragraph.text), None)
            if removable is None:
                return False
            remove_paragraph(removable)
            paragraphs.remove(removable)
            bullet_count -= 1
            bounds = role_range(change.employer, change.role_context)
            if bounds is None:
                return False

        employer_paragraphs = paragraphs[bounds[0]:bounds[1]]
        anchor = next((paragraph for paragraph in employer_paragraphs
                       if change.anchor_text.strip() in paragraph.text), None)
        if anchor is None:
            return False
        new_paragraph = insert_paragraph_after(anchor, change.new_bullet.strip())
        paragraphs.insert(paragraph_index(anchor) + 1, new_paragraph)
        return bullet_count < 4

    def insert_paragraph_after(paragraph: Paragraph, text: str):
        new_element = deepcopy(paragraph._p)
        for child in list(new_element):
            if child.tag != "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr":
                new_element.remove(child)
        paragraph._p.addnext(new_element)
        new_paragraph = Paragraph(new_element, paragraph._parent)

        new_run = OxmlElement("w:r")
        if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
            new_run.append(deepcopy(paragraph.runs[0]._r.rPr))
        text_element = OxmlElement("w:t")
        text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text_element.text = text
        new_run.append(text_element)
        new_element.append(new_run)
        return new_paragraph

    for change in replacements:
        target = change.original_text.strip()
        new_val = change.replacement_text.strip()
        
        if not target:
            continue

        def process_paragraph(p):
            nonlocal applied_count
            if replace_paragraph_text(p, target, new_val):
                applied_count += 1
                return True
            return False

        skill_start, skill_end = range_for_heading("Skills", ("Profile", "Employment History"))
        if "skill" in change.section_name.lower():
            for p in paragraphs[skill_start:skill_end]:
                process_paragraph(p)
        else:
            for p in paragraphs[:employment_start] + paragraphs[employment_end:]:
                process_paragraph(p)

    for change in employment_changes:
        if apply_employment_change(change):
            applied_count += 1

    for section in priority_data.priority_sections:
        is_employment_section = "employment" in section.section_title.lower() or "experience" in section.section_title.lower()
        if not is_employment_section:
            continue

        for item in sorted(section.items, key=lambda priority_item: priority_item.priority_rank):
            if item.action_type.lower().startswith("refram") and item.original_text.strip() and item.content_text.strip():
                target = item.original_text.strip()
                replacement = item.content_text.strip()
                employer = item.employer
                role_context = item.role_context
                bounds = role_range(employer, role_context) if employer else None
                if bounds and any(
                    replace_paragraph_text(paragraph, target, replacement)
                    for paragraph in paragraphs[bounds[0]:bounds[1]]
                ):
                    applied_count += 1
            elif item.action_type.lower().startswith("new") and item.anchor_text.strip() and item.content_text.strip():
                employer = item.employer
                if employer:
                    change = EmploymentHistoryChange(
                        employer=employer,
                        role_context=item.role_context,
                        job_requirement=item.job_requirement_matched,
                        matched_candidate_experience=item.content_text,
                        action_type=item.action_type,
                        anchor_text=item.anchor_text,
                        new_bullet=item.content_text,
                        remove_text=item.remove_text,
                    )
                    if apply_employment_change(change):
                        applied_count += 1
    doc.save(output_docx_path)
    print(f"✓ Applied {applied_count} text updates to the Word structure.")

def print_terminal_changelog(changes: List[BulletChange], priority_data: ReorderAndEnhanceResponse):
    """Outputs clear rationale, text changes, and priority placement reports to terminal."""
    print("\n[4/5] Writing the tailoring and employment-history report...")
    print("\n" + "="*70)
    print("📌 TAILORING CHANGELOG & TEXT REPLACEMENTS")
    print("="*70)

    for i, item in enumerate(changes, 1):
        print(f"\n🔹 CHANGE #{i} [{item.section_name.upper()}]")
        print(f"   • Job Requirement : {item.job_requirement}")
        print(f"   • Your Experience : {item.matched_candidate_experience}")
        print(f"   • Original Text   : \"{item.original_text}\"")
        print(f"   • Updated To      : \"{item.replacement_text}\"")

    print("\n" + "="*70)
    print("📊 PRIORITY RE-ORDERING & SKILL PLACEMENT REPORT")
    print("="*70)

    for sec in priority_data.priority_sections:
        print(f"\n📂 SECTION: {sec.section_title.upper()}")
        sorted_items = sorted(sec.items, key=lambda x: x.priority_rank)
        for item in sorted_items:
            print(f"  [Rank #{item.priority_rank}] ({item.action_type.upper()})")
            print(f"   • Job Priority Match : {item.job_requirement_matched}")
            print(f"   • Positioned Text    : \"{item.content_text}\"")
            if item.original_text:
                print(f"   • Replaces           : \"{item.original_text}\"")
            if item.anchor_text:
                print(f"   • Inserts After      : \"{item.anchor_text}\"")

    print("\n" + "="*70)

def print_employment_history_changelog(changes: List[EmploymentHistoryChange]):
    """Outputs employer-specific changes supported by the verified history file."""
    if not changes:
        return

    print("\n" + "="*70)
    print("💼 EMPLOYMENT HISTORY CHANGES")
    print("="*70)

    for i, change in enumerate(changes, 1):
        print(f"\n🔹 EMPLOYMENT CHANGE #{i} [{change.employer.upper()}]")
        print(f"   • Action            : {change.action_type}")
        print(f"   • Job Requirement   : {change.job_requirement}")
        print(f"   • Verified Evidence : {change.matched_candidate_experience}")
        if change.original_text:
            print(f"   • Original Text     : \"{change.original_text}\"")
        if change.replacement_text:
            print(f"   • Updated To        : \"{change.replacement_text}\"")
        if change.new_bullet:
            print(f"   • New Bullet        : \"{change.new_bullet}\"")

    print("\n" + "="*70)

def run_agent():
    print("=== AI RESUME TAILORING AGENT ===")
    
    job_url = input("Enter the Job Posting URL: ").strip()
    if not job_url:
        print("URL required!")
        return

    master_docx_path = "master_resume.docx"
    employment_history = load_employment_history()

    # 1. Scrape Job Posting
    job_desc, was_scraped = get_job_description(job_url)

    # 2. Read the master resume
    try:
        master_text = extract_text_from_docx(master_docx_path)

        # 3. Generate all tailoring decisions in one Gemini request
        consolidated_response = get_consolidated_tailoring(master_text, job_desc, employment_history)
        meta = {
            "company": consolidated_response.company,
            "title": consolidated_response.title,
        }
        def clean_metadata_value(value, fallback: str) -> str:
            if not isinstance(value, str) or not value.strip():
                value = fallback
            cleaned = re.sub(r"[^\w\-_]", "_", value).strip("_")
            return cleaned or fallback

        meta["clean_company"] = clean_metadata_value(meta["company"], "TargetCompany")
        meta["clean_title"] = clean_metadata_value(meta["title"], "Role")
        file_suffix = f"{meta['clean_company']}_{meta['clean_title']}"

        # Generate unique output file paths to prevent overwriting existing files
        base_docx_path = f"tailored_resume_{file_suffix}.docx"
        output_docx_path = get_unique_filepath(base_docx_path)

        priority_response = ReorderAndEnhanceResponse(
            priority_sections=consolidated_response.priority_sections
        )
    except GeminiRequestError as error:
        print(f"\n❌ {error}")
        print("   The existing resume files were left unchanged.")
        return

    # 4. Output Combined Terminal Report
    print_terminal_changelog(consolidated_response.changes, priority_response)
    print_employment_history_changelog(consolidated_response.employment_history_changes)

    # 5. Apply In-Place Replacements to Word Template
    replace_text_in_docx_inplace(
        master_docx_path,
        consolidated_response.changes,
        consolidated_response.employment_history_changes,
        priority_response,
        output_docx_path,
    )

if __name__ == "__main__":
    run_agent()